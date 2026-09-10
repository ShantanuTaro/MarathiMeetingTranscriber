"""Self-check: run `python test_transcribe.py`. No audio, no network, no key."""
from types import SimpleNamespace as NS

from docx.enum.text import WD_COLOR_INDEX

from transcribe import _group, build_docx, looping, ts


def W(word, prob, start):
    return NS(word=word, probability=prob, start=start, end=start + 0.4)


def S(start, end, text, words=()):
    return NS(start=start, end=end, text=text, words=list(words))


def main():
    assert ts(0) == "00:00"
    assert ts(65) == "01:05"
    assert ts(3725) == "1:02:05"

    segments = [
        S(0.0, 2.0, "नमस्कार सर्वांना",
          [W(" नमस्कार", 0.97, 0.1), W(" सर्वांना", 0.91, 1.0)]),
        S(61.0, 63.0, "आजची अजेंडा meeting",
          [W(" आजची", 0.88, 61.0), W(" अजेंडा", 0.35, 61.5), W(" meeting", 0.92, 62.0)]),
        S(70.0, 71.0, "fallback"),                      # no word timings
    ]

    # named explicitly: the default backend is Sarvam now, and Sarvam's document
    # disowns the confidence summary this block goes on to check.
    doc, flagged = build_docx(segments, threshold=0.6, title="clip.m4a",
                              total_duration=71.0, asr="elevenlabs")

    # low confidence flags on the score; Latin script flags whatever its score
    assert flagged == [("01:01", "अजेंडा", "35% sure"),
                       ("01:02", "meeting", "Latin script")], flagged

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "2 word(s) need a look" in text
    assert "1 below 60% confidence" in text and "1 still in Latin" in text
    assert "00:00  नमस्कार सर्वांना" in text, text
    assert "01:10  fallback" in text, "a segment without word timings must still render"

    meta, review = doc.tables
    assert [r.cells[0].text for r in meta.rows][:2] == ["SOURCE", "DURATION"]
    assert [c.text for c in review.rows[0].cells] == \
        ["#", "TIME", "HEARD AS", "WHY", "CORRECTION"]
    assert [c.text for c in review.rows[2].cells] == \
        ["2", "01:02", "meeting", "Latin script", ""]

    # the checklist comes before the transcript it refers to
    body = list(doc.element.body)
    assert body.index(review._element) < \
        body.index(next(p for p in doc.paragraphs if p.text == "Transcript")._element)

    highlighted = [(r.text.strip(), r.font.highlight_color) for p in doc.paragraphs
                   for r in p.runs if r.font.highlight_color is not None]
    assert highlighted == [("अजेंडा", WD_COLOR_INDEX.YELLOW),
                           ("meeting", WD_COLOR_INDEX.TURQUOISE)], highlighted

    _check_grouping()
    _check_looping()
    _check_scribe()
    _check_sarvam()
    _check_sarvam_split()
    _check_backend_claims()
    _check_unknown_backend()
    _check_billing()
    print("ok")


def _check_grouping():
    """Consecutive speech merges; a real pause or 30s of talk starts a paragraph."""
    close = [S(0, 3, "a"), S(3.2, 6, "b"), S(6.1, 9, "c")]
    assert [len(g) for g in _group(close)] == [3]

    pause = [S(0, 3, "a"), S(9, 12, "b")]                 # 6s gap
    assert [len(g) for g in _group(pause)] == [1, 1]

    # 45s of unbroken speech: the first paragraph closes once it passes 30s
    long_run = [S(i * 5, i * 5 + 4.8, "x") for i in range(9)]
    assert [len(g) for g in _group(long_run)] == [6, 3], \
        [len(g) for g in _group(long_run)]

    assert list(_group([])) == []


def _check_scribe():
    """ElevenLabs' flat token stream -> segments, with spacing folded in and the
    logprob back on CONF_THRESHOLD's scale."""
    from transcribe import _scribe_segments

    def w(text, start, end, kind="word", spk="speaker_0", lp=0.0):
        return {"text": text, "start": start, "end": end, "type": kind,
                "speaker_id": spk, "logprob": lp}

    segs = _scribe_segments({"words": [
        w("नमस्कार", 0.0, 0.5),
        w(" ", 0.5, 0.6, "spacing"),
        w("सर्व", 0.6, 1.0, lp=-2.3),                 # ~10%: below any threshold
        w("[laughter]", 1.0, 1.2, "audio_event"),     # not speech, not a word
        w("मित्रांनो", 3.0, 3.4),                       # 2s gap -> new segment
        w(" ", 3.4, 3.5, "spacing"),
        w("हो", 3.5, 3.8, spk="speaker_1"),           # new voice -> new segment
    ]})

    assert [len(s.words) for s in segs] == [2, 1, 1], segs
    # every run carries its own leading space, first included, as Whisper's do
    assert segs[0].text == " नमस्कार सर्व", segs[0].text
    assert [s.speaker for s in segs] == [0, 0, 1], segs   # ids interned in first-seen order
    assert segs[0].start == 0.0 and segs[0].end == 1.0
    assert abs(segs[0].words[0].probability - 1.0) < 1e-6
    assert 0.09 < segs[0].words[1].probability < 0.11, segs[0].words[1].probability
    assert segs[1].words[0].word.startswith(" "), "runs print back to back"
    assert _scribe_segments({}) == []


def _check_billing():
    """The session meter follows Sarvam's billing, not the job's outcome: charged when
    the backend says it started, never on success, never refunded on a failure."""
    import os
    from pathlib import Path

    import app, transcribe

    app.SESSIONS["t"] = {"asr": "sarvam", "key": "k", "spent": 0.0, "balance": 100.0,
                         "jobs": 0, "rate": 45}
    transcribe.duration = lambda path: 3600.0          # an hour, without an hour of audio
    app.core.duration = transcribe.duration

    def run(fake):
        app.JOBS["j"] = {"status": "queued"}
        app.core.transcribe = fake
        app.run_job("j", Path("nothing.m4a"), "nothing.m4a", "sarvam", "k", "t")
        return app.JOBS["j"]["status"], app.SESSIONS["t"]["spent"]

    # died before the backend started: the audio never moved, so nothing is owed
    def refused(path, on_billed=None, **kw):
        raise RuntimeError("Sarvam 403")
    assert run(refused) == ("error", 0.0), app.SESSIONS["t"]

    # started, then came back empty: billed in full anyway, and it stays billed
    def started_then_failed(path, on_billed=None, **kw):
        on_billed(3600.0)
        raise RuntimeError("empty transcript")
    assert run(started_then_failed) == ("error", 45.0), app.SESSIONS["t"]
    assert app.SESSIONS["t"]["jobs"] == 1
    assert app._wallet("t")["remaining"] == 55.0, app._wallet("t")

    # one token is one backend: a Sarvam key must never ride to ElevenLabs
    assert app.session_key("t", "sarvam") == "k"
    assert app.session_key("t", "elevenlabs") is None

    # ending the session drops the key, and with it any way to read the meter
    app.SESSIONS.pop("t")
    assert app.session_key("t", "sarvam") is None
    assert app._wallet("t") is None
    if not os.getenv("SARVAM_API_KEY"):        # a server-side key outlives the session
        # still listed (the page can take another key) but no longer runnable
        assert {o["id"]: o["ready"] for o in app.options("t")}["sarvam"] is False


def _check_sarvam_split():
    """A recording over Sarvam's per-file cap goes up in parts. Each transcript comes
    back timed from its own zero, so the stitch is where the clock is decided."""
    import transcribe

    def part():
        """What one Sarvam job hands back: timed from zero, speakers numbered from
        zero, and no idea another part exists."""
        return [NS(start=0.0, end=2.0, text=" start", speaker=0,
                   words=[W(" start", 1.0, 0.0)]),
                NS(start=10.0, end=12.0, text=" end", speaker=0,
                   words=[W(" end", 1.0, 10.0)])]

    transcribe.duration = lambda path: 8000.0            # over the 2h cap, cut in half
    transcribe._split_audio = lambda path, secs, into: [("a.m4a", 0.0), ("b.m4a", 3600.0)]
    def one(path, on_progress=None, should_stop=None, key=None, on_billed=None,
            on_stage=None):
        on_stage("waiting")          # the real one names its step; the wrapper adds
        return part()                # which part it belongs to
    transcribe._sarvam_one = one

    stages = []
    segs = transcribe._sarvam("long.m4a", on_stage=stages.append)

    # part two is pushed onto the first part's clock, words along with the segments
    assert [s.start for s in segs] == [0.0, 10.0, 3600.0, 3610.0], [s.start for s in segs]
    assert [s.words[0].start for s in segs] == [0.0, 10.0, 3600.0, 3610.0]
    # "speaker 0" in part two is not the person it was in part one, so it is renumbered
    assert [s.speaker for s in segs] == [0, 0, 1, 1], [s.speaker for s in segs]
    # splitting is a step of its own: re-encoding two hours takes minutes, and
    # without it the card sits on "preparing" with nothing to show.
    assert stages[0] == "splitting", stages
    assert stages[1].endswith("part 1 of 2") and stages[-1].endswith("part 2 of 2")


def _check_backend_claims():
    """A job picks its own backend, so the document's claims must follow the segments
    that built it, not whatever ASR the server happened to start with."""
    from transcribe import model_name

    segs = [S(0.0, 1.0, "meeting", [W(" meeting", 1.0, 0.0)])]
    for asr, named in (("sarvam", True), ("elevenlabs", False)):
        doc, _ = build_docx(segs, threshold=0.6, title="clip.m4a", asr=asr)
        text = "\n".join(p.text for p in doc.paragraphs)
        meta = dict((r.cells[0].text, r.cells[1].text) for r in doc.tables[0].rows)
        assert meta["MODEL"] == model_name(asr), (asr, meta["MODEL"])
        # only Sarvam has to disown the yellow highlight, and only Sarvam may
        assert ("reports no per-word confidence" in text) is named, asr
        assert ("below 60% confidence" in text) is not named, asr


def _check_sarvam():
    """Sarvam's chunk-level output -> segments, with words synthesised so the
    Latin-script flag still has something to walk."""
    from transcribe import _sarvam_segments, reason

    diarized = _sarvam_segments({"diarized_transcript": {"entries": [
        {"transcript": "नमस्कार सर्व", "start_time_seconds": 0.0,
         "end_time_seconds": 2.0, "speaker_id": "0"},
        {"transcript": "budget review आहे", "start_time_seconds": 2.0,
         "end_time_seconds": 5.0, "speaker_id": "1"},
        {"transcript": "   ", "start_time_seconds": 5.0,          # dropped
         "end_time_seconds": 5.1, "speaker_id": "1"},
    ]}})
    assert [s.speaker for s in diarized] == [0, 1], diarized
    assert [len(s.words) for s in diarized] == [2, 3]
    # timings interpolated across the chunk, so a review row lands in the right minute
    assert diarized[0].words[1].start == 1.0
    assert diarized[1].words[-1].end == 5.0
    assert all(w.word.startswith(" ") for s in diarized for w in s.words)

    # no confidence from Sarvam: only Latin script can flag a word, never "x% sure"
    flags = [reason(w, 0.60) for s in diarized for w in s.words]
    assert flags == [None, None, "Latin script", "Latin script", None], flags

    # with_diarization refused -> the three parallel lists. The live API keys the
    # list "words" and the docs call it "chunks"; both are sentences, both parse.
    plain = _sarvam_segments({"timestamps": {
        "words": ["एक दोन", "तीन"],
        "start_time_seconds": [0.0, 3.0], "end_time_seconds": [2.0, 4.0]}})
    assert plain == _sarvam_segments({"timestamps": {
        "chunks": ["एक दोन", "तीन"],
        "start_time_seconds": [0.0, 3.0], "end_time_seconds": [2.0, 4.0]}})
    assert [s.speaker for s in plain] == [None, None]
    assert plain[1].text == " तीन", plain[1].text

    # with_timestamps refused too -> one segment rather than an empty document
    assert len(_sarvam_segments({"transcript": "काहीतरी"})) == 1
    assert _sarvam_segments({}) == []


def _check_unknown_backend():
    """The local path is gone. Asking for it must say so, not silently do something
    else with an hour of audio and somebody's money."""
    import transcribe
    try:
        transcribe.transcribe("clip.m4a", asr="mlx")
    except ValueError as e:
        assert "mlx" in str(e) and "sarvam" in str(e), e
    else:
        raise AssertionError("an unknown backend must not be accepted")


def _check_looping():
    assert looping("प्रदावादा " * 40)                    # the hall-audio failure mode
    assert not looping("ग्रामसभा साथी सचीव जवा आमि प्रस्थिता सेव तेरा सचीवाची निवड सर्पंज करूँ सकता")
    assert not looping("छोटा")                          # too short to judge


if __name__ == "__main__":
    main()
